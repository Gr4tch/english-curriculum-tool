"""
doc_exporter.py
Standalone export engine for 8th Grade U.S. History Curriculum materials.
Handles document styling, native table generation, metadata pill formatting,
compact page fitting, DOCX byte generation, and Google Drive / Google Docs export.
"""

import os
import re
import json
import logging
from pathlib import Path
from typing import Dict, Any, Tuple, Optional, List

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
import googleapiclient.discovery
from googleapiclient.http import MediaFileUpload

logger = logging.getLogger(__name__)

# Standard locations for credentials and tokens
TOKEN_FILE_PATHS = [
    Path(r"C:\Users\Jared\.gemini\antigravity\plugin_data\gdrive\token.json"),
    Path("output/gdrive_token.json"),
    Path(".gdrive_token.json"),
]
ROOT_FOLDER_NAME = "#US History (English) Curriculum Materials"
TEMP_DOCS_DIR = Path("output/google_docs_temp")


# ─────────────────────────────────────────────────────────────────────────────
# 1. AUTHENTICATION & CREDENTIALS MANAGEMENT
# ─────────────────────────────────────────────────────────────────────────────

def get_token_file() -> Optional[Path]:
    """Finds the existing Google Drive OAuth token file."""
    for path in TOKEN_FILE_PATHS:
        if path.exists():
            return path
    return None


def get_gdrive_service():
    """
    Builds and returns an authenticated Google Drive API v3 service.
    Automatically refreshes the token if expired and saves the refreshed state.
    Supports local token files as well as GDRIVE_TOKEN_JSON from Streamlit secrets / environment.
    """
    token_data = None
    token_path = None

    # Check Streamlit Cloud secrets or environment variable first
    try:
        import streamlit as st
        if "GDRIVE_TOKEN_JSON" in st.secrets:
            sec_val = st.secrets["GDRIVE_TOKEN_JSON"]
            if isinstance(sec_val, (dict, list)):
                token_data = dict(sec_val)
            elif isinstance(sec_val, str):
                token_data = json.loads(sec_val.strip())
    except Exception as e:
        logger.warning(f"Could not load GDRIVE_TOKEN_JSON from st.secrets: {e}")

    if not token_data and "GDRIVE_TOKEN_JSON" in os.environ:
        try:
            token_data = json.loads(os.environ["GDRIVE_TOKEN_JSON"].strip())
        except Exception as e:
            logger.warning(f"Could not load GDRIVE_TOKEN_JSON from os.environ: {e}")

    if not token_data:
        token_path = get_token_file()
        if not token_path or not token_path.exists():
            raise FileNotFoundError(
                "Google Drive credentials not found. To enable Google Drive exports on Streamlit Cloud, "
                "please configure your GDRIVE_TOKEN_JSON secret in Streamlit Settings -> Secrets."
            )
        token_data = json.loads(token_path.read_text(encoding="utf-8"))

    tok = token_data.get("token", {})
    if isinstance(tok, str):
        try:
            tok = json.loads(tok)
        except Exception:
            pass

    access_token = tok.get("access_token") if (isinstance(tok, dict) or hasattr(tok, "get")) else token_data.get("access_token")
    refresh_token = tok.get("refresh_token") if (isinstance(tok, dict) or hasattr(tok, "get")) else token_data.get("refresh_token")
    client_id = token_data.get("client_id")
    client_secret = token_data.get("client_secret")
    scopes = token_data.get("scopes", ["https://www.googleapis.com/auth/drive"])
    if isinstance(scopes, str):
        scopes = [s.strip() for s in scopes.split(",")]

    if not access_token or not refresh_token:
        raise ValueError("Invalid Google Drive token structure: missing access or refresh token.")

    creds = Credentials(
        token=access_token,
        refresh_token=refresh_token,
        token_uri="https://oauth2.googleapis.com/token",
        client_id=client_id,
        client_secret=client_secret,
        scopes=scopes,
    )

    if creds.expired or not creds.valid:
        req = Request()
        creds.refresh(req)
        if isinstance(token_data.get("token"), dict):
            token_data["token"]["access_token"] = creds.token
            if creds.refresh_token:
                token_data["token"]["refresh_token"] = creds.refresh_token
        else:
            token_data["access_token"] = creds.token
            if creds.refresh_token:
                token_data["refresh_token"] = creds.refresh_token
        if token_path:
            try:
                token_path.write_text(json.dumps(token_data, indent=2), encoding="utf-8")
            except Exception as e:
                logger.warning(f"Could not persist refreshed token back to {token_path}: {e}")

    return googleapiclient.discovery.build("drive", "v3", credentials=creds)


def get_connected_user_info(service) -> Dict[str, str]:
    """Retrieves display name and email address of the authenticated Google account."""
    try:
        about = service.about().get(fields="user(displayName, emailAddress)").execute()
        user = about.get("user", {})
        return {
            "name": user.get("displayName", "Google User"),
            "email": user.get("emailAddress", ""),
        }
    except Exception as e:
        logger.warning(f"Could not retrieve Google user info: {e}")
        return {"name": "Google Drive User", "email": ""}


def grant_file_permissions(
    service,
    file_id: str,
    role: str = "writer",
    share_public: bool = True,
    additional_email: Optional[str] = None,
) -> Dict[str, bool]:
    """
    Grants link sharing permissions ('anyone' with 'writer' or 'reader') so the colleague can
    access the parent folder via a link without needing his email address upfront.
    Optionally grants explicit user permission if an email is provided later.
    """
    status = {"public": False, "user": False}

    if share_public:
        try:
            service.permissions().create(
                fileId=file_id,
                body={"type": "anyone", "role": role},
                fields="id",
            ).execute()
            status["public"] = True
        except Exception as e:
            logger.warning(f"Failed to grant 'anyone' {role} permission on {file_id}: {e}")
            if role == "writer":
                try:
                    service.permissions().create(
                        fileId=file_id,
                        body={"type": "anyone", "role": "reader"},
                        fields="id",
                    ).execute()
                    status["public"] = True
                except Exception as e2:
                    logger.warning(f"Failed to grant 'anyone' reader permission on {file_id}: {e2}")

    if additional_email and "@" in additional_email:
        clean_email = additional_email.strip()
        try:
            service.permissions().create(
                fileId=file_id,
                body={"type": "user", "role": "writer", "emailAddress": clean_email},
                sendNotificationEmail=False,
                fields="id",
            ).execute()
            status["user"] = True
        except Exception as e:
            logger.debug(f"Explicit user permission on {file_id} for {clean_email}: {e}")

    return status


# ─────────────────────────────────────────────────────────────────────────────
# 2. DRIVE FOLDER HIERARCHY MANAGEMENT
# ─────────────────────────────────────────────────────────────────────────────

def get_or_create_folder(
    service,
    folder_name: str,
    parent_id: Optional[str] = None,
    additional_email: Optional[str] = None,
) -> Tuple[str, str]:
    """
    Finds a folder by name under the specified parent (or root).
    If it does not exist, creates it and applies link-sharing permissions.
    Returns (folder_id, webViewLink).
    """
    safe_name = folder_name.replace("'", "\\'")
    if parent_id:
        query = (
            f"name = '{safe_name}' and '{parent_id}' in parents "
            f"and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
        )
    else:
        query = (
            f"name = '{safe_name}' and mimeType = 'application/vnd.google-apps.folder' "
            f"and trashed = false"
        )

    res = service.files().list(q=query, fields="files(id, name, webViewLink)").execute()
    folders = res.get("files", [])

    if folders:
        fid = folders[0]["id"]
        link = folders[0].get("webViewLink", f"https://drive.google.com/drive/folders/{fid}")
        return fid, link

    metadata: Dict[str, Any] = {
        "name": folder_name,
        "mimeType": "application/vnd.google-apps.folder",
    }
    if parent_id:
        metadata["parents"] = [parent_id]

    created = service.files().create(body=metadata, fields="id, name, webViewLink").execute()
    folder_id = created["id"]
    folder_link = created.get("webViewLink", f"https://drive.google.com/drive/folders/{folder_id}")
    grant_file_permissions(
        service,
        folder_id,
        role="writer",
        share_public=True,
        additional_email=additional_email,
    )
    return folder_id, folder_link


def get_target_chapter_folder(
    service,
    chapter_num: int,
    chapter_title: str,
    additional_email: Optional[str] = None,
) -> Tuple[str, str, str]:
    """
    Ensures the chapter folder hierarchy:
      #US History (English) Curriculum Materials / Chp XX - [Title]
    Returns (chapter_folder_id, chapter_folder_path_display, chapter_folder_url).
    """
    root_id, _ = get_or_create_folder(service, ROOT_FOLDER_NAME, additional_email=additional_email)
    clean_ch_title = chapter_title.strip()
    chapter_folder_name = f"Chp {chapter_num:02d} - {clean_ch_title}"
    chapter_id, chapter_url = get_or_create_folder(
        service,
        chapter_folder_name,
        parent_id=root_id,
        additional_email=additional_email,
    )
    path_display = f"{ROOT_FOLDER_NAME} > {chapter_folder_name}"
    return chapter_id, path_display, chapter_url


def get_target_section_folder(
    service,
    chapter_num: int,
    chapter_title: str,
    section_code: str,
    section_title: str,
    additional_email: Optional[str] = None,
) -> Tuple[str, str, str]:
    """
    Ensures the folder hierarchy:
      #US History (English) Curriculum Materials / Chp XX - [Title] / Sec. X.X - [Title]
    Returns (section_folder_id, section_folder_path_display, section_folder_url).
    """
    chapter_id, ch_path, _ = get_target_chapter_folder(
        service,
        chapter_num=chapter_num,
        chapter_title=chapter_title,
        additional_email=additional_email,
    )
    clean_sec_title = section_title.strip()
    section_folder_name = f"Sec. {section_code} - {clean_sec_title}"
    section_id, section_url = get_or_create_folder(
        service,
        section_folder_name,
        parent_id=chapter_id,
        additional_email=additional_email,
    )
    path_display = f"{ch_path} > {section_folder_name}"
    return section_id, path_display, section_url


# ─────────────────────────────────────────────────────────────────────────────
# 3. VERSIONING & DUPLICATE FILE HANDLING
# ─────────────────────────────────────────────────────────────────────────────

def get_versioned_filename(
    service,
    parent_folder_id: str,
    base_name: str,
    replace_existing: bool = True,
) -> str:
    """
    Checks if a document with base_name exists in the destination folder.
    If replace_existing is True (default), trashes any older copies so only the clean,
    single latest file remains with no '(v2)' or '(v3)' clutter.
    """
    safe_base = base_name.replace("'", "\\'")
    query = f"'{parent_folder_id}' in parents and trashed = false and name contains '{safe_base}'"
    res = service.files().list(q=query, fields="files(id, name)").execute()
    existing_files = res.get("files", [])

    if replace_existing:
        for f in existing_files:
            try:
                service.files().update(fileId=f["id"], body={"trashed": True}).execute()
            except Exception:
                pass
        return base_name

    existing_names = [f["name"] for f in existing_files]
    if not existing_names or base_name not in existing_names:
        return base_name

    highest_ver = 1
    pattern = re.compile(rf"^{re.escape(base_name)}(?:\s*\(v(\d+)\))?$")
    for fname in existing_names:
        m = pattern.match(fname)
        if m:
            ver = int(m.group(1)) if m.group(1) else 1
            if ver > highest_ver:
                highest_ver = ver

    return f"{base_name} (v{highest_ver + 1})"


# ─────────────────────────────────────────────────────────────────────────────
# 4. LOCKED-IN VIBRANT-LOVELACE DOCX BUILDER (ADAPTED FOR ENGLISH)
# ─────────────────────────────────────────────────────────────────────────────

def set_cell_border(cell, color: str = "CCCCCC", sz: str = "4"):
    """Sets borders for a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(existing)
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_cell_v_align(cell, val: str = "top"):
    """Sets vertical alignment for a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(existing)
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{val}"/>')
    tcPr.append(vAlign)


def set_row_height(row, height_twips: int, rule: str = "atLeast"):
    """Sets row height in twips (1 inch = 1440 twips)."""
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    for existing in trPr.findall(qn("w:trHeight")):
        trPr.remove(existing)
    trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{height_twips}" w:hRule="{rule}"/>')
    trPr.append(trHeight)


def add_metadata_pills_table(
    doc,
    pills: List[Tuple[str, str, str]],
    width_inches: float = 7.0,
):
    """
    Renders metadata (Day, Standard, Points, Edition) inside a single-row, multi-column table
    with NO borders (all borders set to 'none'), subtle background shading (e.g. #F2F4F8),
    and bold colored text.
    """
    num_pills = len(pills)
    if num_pills == 0:
        return None
    num_cols = num_pills * 2 - 1
    table = doc.add_table(rows=1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    pill_w = [1.3, 2.3, 1.2, 1.75] if num_pills == 4 else [width_inches / num_pills] * num_pills
    spacer_w = 0.15

    col_widths = []
    for i in range(num_pills):
        col_widths.append(Inches(pill_w[i] if i < len(pill_w) else 1.5))
        if i < num_pills - 1:
            col_widths.append(Inches(spacer_w))

    for c_idx, w in enumerate(col_widths):
        table.columns[c_idx].width = w

    for i, (text, bg_hex, text_hex) in enumerate(pills):
        c_idx = i * 2
        cell = table.cell(0, c_idx)
        cell.width = col_widths[c_idx]

        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(borders)

        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="50" w:type="dxa"/>
                <w:left w:w="80" w:type="dxa"/>
                <w:bottom w:w="50" w:type="dxa"/>
                <w:right w:w="80" w:type="dxa"/>
            </w:tcMar>
        ''')
        cell._tc.get_or_add_tcPr().append(tcMar)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(text_hex)

    for i in range(num_pills - 1):
        s_idx = i * 2 + 1
        s_cell = table.cell(0, s_idx)
        s_cell.width = col_widths[s_idx]
        s_borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        s_cell._tc.get_or_add_tcPr().append(s_borders)
        s_cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        s_cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0
    sp.paragraph_format.keep_with_next = True
    return table


def add_student_response_table(
    doc,
    num_lines: int = 3,
    width_inches: float = 7.0,
    is_note_taking: bool = False,
):
    """
    Generates clean 1-column handwriting response tables where only the bottom border is visible
    (border-bottom: 0.5pt solid #CCCCCC), with cell height set to 0.3 inches per line (432 twips).
    Avoids raw underscore characters ('____').
    """
    table = doc.add_table(rows=num_lines, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(width_inches)

    border_color = "94A3B8" if is_note_taking else "CCCCCC"
    border_val = "dotted" if is_note_taking else "single"

    for r_idx in range(num_lines):
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="432" w:hRule="atLeast"/>')
        trPr.append(trHeight)

        cell = row.cells[0]
        cell.width = Inches(width_inches)

        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="{border_val}" w:sz="4" w:space="0" w:color="{border_color}"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(borders)

        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="40" w:type="dxa"/>
                <w:left w:w="60" w:type="dxa"/>
                <w:bottom w:w="40" w:type="dxa"/>
                <w:right w:w="60" w:type="dxa"/>
            </w:tcMar>
        ''')
        cell._tc.get_or_add_tcPr().append(tcMar)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.15
    return table


def add_styled_paragraph(
    doc,
    text: str,
    style: str = "Normal",
    space_after: int = 2,
    space_before: int = 2,
    is_quote: bool = False,
    keep_with_next: bool = False,
):
    """Adds a paragraph with inline bold/italic parsing, line_spacing=1.15, and custom spacing."""
    text = text.replace(r"$\rightarrow$", "→").replace(r"\rightarrow", "→").replace("$", "")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if is_quote:
        p.paragraph_format.left_indent = Inches(0.35)

    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = p.add_run(tok[2:-2])
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            r.italic = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
        else:
            r = p.add_run(tok)
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            if is_quote:
                r.italic = True
                r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    return p


def _populate_formatted_cell(
    cell,
    raw_text: str,
    is_header: bool = False,
    font_size_bold: float = 10.0,
    font_size_reg: float = 9.5,
    center_align: bool = False,
    prepend_blank_line: bool = False,
):
    """Populates a table cell with support for <br> line breaks and **bold** / *italic* inline tokens."""
    raw_text = raw_text.replace(r"$\rightarrow$", "→").replace(r"\rightarrow", "→").replace("$", "")
    cell.text = ""
    sublines = re.split(r'<br\s*/?>', raw_text, flags=re.IGNORECASE)
    p_offset = 0
    if prepend_blank_line:
        p_blank = cell.paragraphs[0]
        p_blank.paragraph_format.space_before = Pt(1)
        p_blank.paragraph_format.space_after = Pt(1)
        p_offset = 1

    for idx, sline in enumerate(sublines):
        clean_sline = sline.strip()
        # Inside a table cell, render pure underscore/dot lines as clean border lines (avoiding raw characters)
        if re.match(r'^_{6,}$', clean_sline):
            p = cell.paragraphs[0] if (idx == 0 and p_offset == 0) else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            continue
        elif re.match(r'^\.{6,}$', clean_sline):
            p = cell.paragraphs[0] if (idx == 0 and p_offset == 0) else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="dotted" w:sz="4" w:space="1" w:color="94A3B8"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            continue
        elif re.search(r'_{6,}', clean_sline):
            clean_sline = re.sub(r'_{6,}', "_" * 20, clean_sline)
        elif re.search(r'\.{6,}', clean_sline):
            clean_sline = re.sub(r'\.{6,}', "." * 18, clean_sline)

        p = cell.paragraphs[0] if (idx == 0 and p_offset == 0) else cell.add_paragraph()
        if center_align:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1 if idx > 0 else 2)
        p.paragraph_format.space_after = Pt(2)
        tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', clean_sline)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                r = p.add_run(tok[2:-2])
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(font_size_bold)
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.italic = True
                r.font.name = "Calibri"
                r.font.size = Pt(font_size_reg)
            else:
                r = p.add_run(tok)
                r.font.name = "Calibri"
                r.font.size = Pt(font_size_bold if is_header else font_size_reg)
                if is_header:
                    r.bold = True


def build_docx_file(
    markdown_text: str,
    doc_title: str,
    is_student_handout: bool,
    output_path: Path,
    display_title: Optional[str] = None,
    display_subtitle: Optional[str] = None,
    compact_mode: bool = True,
    omit_auto_title: bool = True,
    include_top_student_header: bool = True,
    metadata_pills: Optional[List[Tuple[str, str, str]]] = None,
) -> Path:
    """
    Generates a publication-grade, print-ready .docx file aligning with word-processing layout standards:
    - 0.75" margins across all sections (7.0" printable width).
    - Document-wide Normal style with line_spacing=1.15, space_before=Pt(2), space_after=Pt(2).
    - Top-right student header: 'Name: ____________________   Date: __________   Period: _____'.
    - Single-row, multi-column metadata pills table with NO borders (all 'none') and subtle shading.
    - Handwriting sections: 1-column response tables with border-bottom: 0.5pt solid #CCCCCC and 0.3-inch row heights.
    - Task headers and prompt directions with keep_with_next=True.
    - Shaded #F8FAFC callout cards with 3pt #1E3A8A left accent bars for blockquotes.
    - Hard page break before Teacher Plan / Answer Key.
    - Simple right-aligned footer score notation: 'Sec X.X   _____/100'.
    """
    doc = Document()

    # Document-wide Normal style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    normal_style.paragraph_format.space_before = Pt(2)
    normal_style.paragraph_format.space_after = Pt(2)
    normal_style.paragraph_format.line_spacing = 1.15

    current_section = doc.sections[0]
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
    current_section.footer.is_linked_to_previous = False
    if current_section.footer.paragraphs:
        current_section.footer.paragraphs[0].text = ""
        current_section.footer.paragraphs[0].paragraph_format.space_before = Pt(0)
        current_section.footer.paragraphs[0].paragraph_format.space_after = Pt(0)
        current_section.footer.paragraphs[0].paragraph_format.line_spacing = Pt(1)

    def _add_english_student_header_line():
        hdr_p = doc.add_paragraph()
        hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hdr_p.paragraph_format.space_before = Pt(0)
        hdr_p.paragraph_format.space_after = Pt(4)
        hdr_p.paragraph_format.line_spacing = 1.15
        hdr_p.paragraph_format.keep_with_next = True
        run_hdr = hdr_p.add_run("Name: ____________________   Date: __________   Period: _____")
        run_hdr.font.name = "Calibri"
        run_hdr.font.size = Pt(10)
        run_hdr.font.bold = True
        run_hdr.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    if is_student_handout and include_top_student_header:
        _add_english_student_header_line()

    if not omit_auto_title:
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(2 if display_subtitle else 8)
        title_p.paragraph_format.line_spacing = 1.15
        title_p.paragraph_format.keep_with_next = True
        main_heading = display_title if display_title else doc_title
        run_title = title_p.add_run(main_heading)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(16.0)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        if display_subtitle:
            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(0)
            sub_p.paragraph_format.space_after = Pt(8)
            sub_p.paragraph_format.line_spacing = 1.15
            sub_p.paragraph_format.keep_with_next = True
            run_sub = sub_p.add_run(display_subtitle)
            run_sub.font.name = "Calibri"
            run_sub.font.size = Pt(11)
            run_sub.font.bold = True
            run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    lines = markdown_text.splitlines()
    in_table = False
    table_lines: List[str] = []
    last_prompt_line = ""
    pending_page_break = False

    def _apply_pending_page_break(p):
        nonlocal pending_page_break
        if pending_page_break:
            p.paragraph_format.page_break_before = True
            pending_page_break = False
        return p

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        parsed_rows = []
        for tl in table_lines:
            if not tl.strip() or re.match(r'^\s*\|?\s*:?-+:?\s*\|', tl):
                continue
            cells = [c.strip() for c in tl.strip().strip("|").split("|")]
            if any(cells):
                parsed_rows.append(cells)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"

            header_row_str = " ".join(parsed_rows[0]).lower()
            is_cover_tracker_table = (
                num_cols in (3, 4)
                and ("study guide" in header_row_str or "score / stamp" in header_row_str or "teacher stamp" in header_row_str)
            )
            is_vocab_table = (
                num_cols == 3
                and ("key term" in header_row_str or "official textbook definition" in header_row_str or "brief note" in header_row_str)
            )
            is_cause_effect_table = (
                num_cols == 2
                and ("historical cause" in header_row_str or "historical effect" in header_row_str)
            )
            is_perspectives_table = (
                num_cols == 2
                and ("perspective" in header_row_str or "group" in header_row_str)
            )
            is_chronological_table = (
                num_cols == 2
                and ("step & subheading" in header_row_str or "locator" in header_row_str or "chronological" in header_row_str)
            )
            is_standards_table = (
                num_cols == 2
                and "standard category" in header_row_str
            )
            is_box_bullets_table = (
                num_cols == 1
                and "box & bullets" in header_row_str
            )

            # Proportional column widths across 7.0-inch printable width
            if is_cover_tracker_table and num_cols == 3:
                col_widths = [Inches(2.0), Inches(3.8), Inches(1.2)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_vocab_table:
                col_widths = [Inches(1.7), Inches(3.3), Inches(2.0)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_cause_effect_table:
                col_widths = [Inches(3.5), Inches(3.5)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_perspectives_table:
                col_widths = [Inches(2.6), Inches(4.4)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_chronological_table:
                col_widths = [Inches(2.5), Inches(4.5)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_standards_table:
                col_widths = [Inches(2.4), Inches(4.6)]
                table.autofit = False
                for c_idx, cw in enumerate(col_widths):
                    table.columns[c_idx].width = cw
            elif is_box_bullets_table:
                col_widths = [Inches(7.0)]
                table.autofit = False
                table.columns[0].width = Inches(7.0)
            else:
                col_widths = [Inches(7.0 / num_cols)] * num_cols

            for r_idx, row_data in enumerate(parsed_rows):
                if is_cover_tracker_table and r_idx > 0:
                    set_row_height(table.rows[r_idx], 400, "atLeast")
                elif is_vocab_table and r_idx > 0:
                    set_row_height(table.rows[r_idx], 440, "atLeast")
                elif is_student_handout and r_idx > 0 and not is_standards_table:
                    set_row_height(table.rows[r_idx], 480, "atLeast")

                for c_idx in range(num_cols):
                    cell = table.cell(r_idx, c_idx)
                    if c_idx < len(col_widths):
                        cell.width = col_widths[c_idx]
                    cell_text = row_data[c_idx] if c_idx < len(row_data) else ""

                    if is_standards_table:
                        _populate_formatted_cell(
                            cell,
                            cell_text,
                            is_header=(r_idx == 0),
                            font_size_bold=(9.5 if r_idx == 0 else 8.5),
                            font_size_reg=8.5,
                            center_align=False,
                        )
                    elif is_cause_effect_table:
                        _populate_formatted_cell(
                            cell,
                            cell_text,
                            is_header=(r_idx == 0),
                            font_size_bold=(11.0 if r_idx == 0 else 10.0),
                            font_size_reg=(11.0 if r_idx == 0 else 9.5),
                            center_align=(r_idx == 0),
                        )
                    else:
                        _populate_formatted_cell(
                            cell,
                            cell_text,
                            is_header=(r_idx == 0),
                            font_size_bold=(10.5 if r_idx == 0 else 9.5),
                            font_size_reg=9.5,
                            center_align=(is_cover_tracker_table and c_idx == num_cols - 1),
                        )

                    set_cell_border(cell, color=("CBD5E1" if is_standards_table else "CCCCCC"))
                    set_cell_v_align(cell, "top")
                    tcMar = parse_xml(r'''
                        <w:tcMar {} >
                            <w:top w:w="80" w:type="dxa"/>
                            <w:left w:w="120" w:type="dxa"/>
                            <w:bottom w:w="80" w:type="dxa"/>
                            <w:right w:w="120" w:type="dxa"/>
                        </w:tcMar>
                    '''.format(nsdecls("w")))
                    cell._tc.get_or_add_tcPr().append(tcMar)
                    if r_idx == 0:
                        shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls("w")))
                        cell._tc.get_or_add_tcPr().append(shading)

            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(3)
            spacer.paragraph_format.line_spacing = 1.15

        table_lines = []
        in_table = False

    line_idx = 0
    while line_idx < len(lines):
        line = lines[line_idx]
        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            line_idx += 1
            continue
        elif in_table:
            flush_table()

        if not stripped:
            line_idx += 1
            continue

        # Hard page break before Teacher Answer Key / Teacher Plan
        is_teacher_key_start = any(k in stripped for k in [
            "Teacher Plan & Clipboard Key",
            "Teacher Answer Key",
            "Master Teacher Plan",
        ])
        if is_teacher_key_start and len(doc.paragraphs) > 1:
            current_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            current_section.top_margin = Inches(0.75)
            current_section.bottom_margin = Inches(0.75)
            current_section.left_margin = Inches(0.75)
            current_section.right_margin = Inches(0.75)
            current_section.footer.is_linked_to_previous = False
            for fp in current_section.footer.paragraphs:
                fp.text = ""

        # Special Token 1: Explicit Page Break with Isolated Section Footer
        if stripped == "[PAGE_BREAK]":
            current_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            current_section.top_margin = Inches(0.75)
            current_section.bottom_margin = Inches(0.75)
            current_section.left_margin = Inches(0.75)
            current_section.right_margin = Inches(0.75)
            current_section.footer.is_linked_to_previous = False
            for fp in current_section.footer.paragraphs:
                fp.text = ""
                fp.paragraph_format.space_before = Pt(0)
                fp.paragraph_format.space_after = Pt(0)
                fp.paragraph_format.line_spacing = Pt(1)
            pending_page_break = False
            last_prompt_line = ""
            line_idx += 1
            continue

        # Special Token 2: Centered Section Heading (e.g., Cover Tracker Table Heading)
        if stripped.startswith("[CENTER_HEADING:") and stripped.endswith("]"):
            ch_text = stripped[len("[CENTER_HEADING:"):-1].strip()
            p = _apply_pending_page_break(doc.add_paragraph())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(ch_text)
            r.font.name = "Calibri"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            line_idx += 1
            continue

        # Special Token 3: Simple, Clean Footer Score Notation (e.g. "Sec 1.1   _____/100")
        if stripped.startswith("[GRADE_BOX:") and stripped.endswith("]"):
            box_text = stripped[len("[GRADE_BOX:"):-1].strip()
            fp = current_section.footer.paragraphs[0]
            fp.text = ""
            fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            fp.paragraph_format.space_before = Pt(0)
            fp.paragraph_format.space_after = Pt(0)
            fr = fp.add_run(box_text)
            fr.font.name = "Calibri"
            fr.font.size = Pt(10.0)
            fr.font.bold = True
            fr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            line_idx += 1
            continue

        # Skip duplicate student header line if already rendered
        if is_student_handout and include_top_student_header and "**Name:**" in stripped and "_____" in stripped:
            line_idx += 1
            continue

        # Section Title with Metadata Pills Table
        if stripped.startswith("## Section "):
            m = re.match(r'^## Section (\d+\.?\d*):?\s*([^(]+)\s*(?:\(Textbook pp\.?\s*([^)]+)\))?(.*)', stripped)
            sec_num = m.group(1) if m else "1.1"
            sec_title = m.group(2).strip() if m else "Section"
            pp_str = m.group(3).strip() if (m and m.group(3)) else ""
            is_scaf = "Scaffolded" in stripped

            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True

            disp_pp = f" (Textbook pp. {pp_str})" if pp_str else ""
            r = p.add_run(f"Section {sec_num}: {sec_title}{disp_pp}")
            r.font.name = "Calibri"
            r.font.size = Pt(13.0)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

            # Metadata Pills Table: Single-row, multi-column table with NO borders and subtle background shading
            if metadata_pills:
                pills_data = metadata_pills
            else:
                day_num = sec_num.split(".")[-1]
                pills_data = [
                    (f"DAY {day_num} • SEC {sec_num}", "F2F4F8", "1E3A8A"),
                    ("CCSS RH.6-8.2 (Central Idea)", "EFF6FF", "1D4ED8"),
                    ("100 POINTS", "F1F5F9", "334155"),
                    ("SCAFFOLDED SUPPORT" if is_scaf else "STUDENT COPY", "FEF3C7" if is_scaf else "F0FDF4", "B45309" if is_scaf else "15803D"),
                ]
            add_metadata_pills_table(doc, pills_data, width_inches=7.0)
            last_prompt_line = stripped
            line_idx += 1
            continue

        # Headings
        if stripped.startswith("#### "):
            h_text = stripped.lstrip("#").strip()
            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            last_prompt_line = h_text
            line_idx += 1
            continue
        elif stripped.startswith("### "):
            h_text = stripped.lstrip("#").strip()
            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            last_prompt_line = h_text
            line_idx += 1
            continue
        elif stripped.startswith("## "):
            h_text = stripped.lstrip("#").strip()
            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(13.0)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            last_prompt_line = h_text
            line_idx += 1
            continue
        elif stripped.startswith("# "):
            h_text = stripped.lstrip("#").strip()
            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(15.0)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            last_prompt_line = h_text
            line_idx += 1
            continue

        # Blockquote / Callout (Render as a shaded card with Academic Navy left accent bar)
        elif stripped.startswith("> "):
            q_text = stripped[2:].strip()
            q_table = doc.add_table(rows=1, cols=1)
            q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            q_cell = q_table.cell(0, 0)
            q_cell.width = Inches(7.0)
            shading = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls("w")))
            q_cell._tc.get_or_add_tcPr().append(shading)
            borders = parse_xml(r'''
                <w:tcBorders {} >
                    <w:top w:val="none"/>
                    <w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/>
                    <w:bottom w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            '''.format(nsdecls("w")))
            q_cell._tc.get_or_add_tcPr().append(borders)
            tcMar = parse_xml(r'''
                <w:tcMar {} >
                    <w:top w:w="80" w:type="dxa"/>
                    <w:left w:w="140" w:type="dxa"/>
                    <w:bottom w:w="80" w:type="dxa"/>
                    <w:right w:w="140" w:type="dxa"/>
                </w:tcMar>
            '''.format(nsdecls("w")))
            q_cell._tc.get_or_add_tcPr().append(tcMar)
            _populate_formatted_cell(q_cell, q_text, is_header=False, font_size_bold=10.0, font_size_reg=9.5)
            sp_after = doc.add_paragraph()
            sp_after.paragraph_format.space_before = Pt(0)
            sp_after.paragraph_format.space_after = Pt(4)
            sp_after.paragraph_format.line_spacing = 1.15
            last_prompt_line = q_text
            line_idx += 1
            continue

        # Handwriting Response Lines: Clean 1-column response table with border-bottom: 0.5pt solid #CCCCCC (0.3" height per line)
        elif re.search(r'^[\._]{15,}$', stripped):
            count = 1
            is_note = "...." in stripped
            while line_idx + 1 < len(lines) and re.search(r'^[\._]{15,}$', lines[line_idx + 1].strip()):
                line_idx += 1
                count += 1
                if "...." in lines[line_idx].strip():
                    is_note = True
            add_student_response_table(doc, num_lines=count, width_inches=7.0, is_note_taking=is_note)
            line_idx += 1
            continue

        # Divider line
        elif stripped in ["---", "***", "___"]:
            p = _apply_pending_page_break(doc.add_paragraph())
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run("_________________________________________________________________________________")
            r.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            line_idx += 1
            continue

        # Normal paragraph or prompt item
        else:
            is_hdr_or_prompt = (
                stripped.startswith("**Part ")
                or stripped.startswith("**Directions")
                or stripped.startswith("**Central Idea")
                or "Directions" in stripped
            )
            p = add_styled_paragraph(
                doc,
                stripped,
                space_after=2,
                space_before=2,
                keep_with_next=is_hdr_or_prompt,
            )
            _apply_pending_page_break(p)
            last_prompt_line = stripped
            line_idx += 1

    if in_table:
        flush_table()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# 5. UPLOAD TO GOOGLE DRIVE (GOOGLE DOCS & SEARCHABLE 2-PAGE SECTION PDFS)
# ─────────────────────────────────────────────────────────────────────────────

def upload_docx_as_google_doc(
    service,
    docx_path: Path,
    doc_name: str,
    parent_folder_id: str,
    additional_email: Optional[str] = None,
) -> Dict[str, str]:
    """
    Uploads a local .docx file with mimeType='application/vnd.google-apps.document'.
    Google Drive automatically converts it into a native Google Doc and applies link sharing.
    """
    file_metadata = {
        "name": doc_name,
        "parents": [parent_folder_id],
        "mimeType": "application/vnd.google-apps.document",
    }
    media = MediaFileUpload(
        str(docx_path),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        resumable=True,
    )

    created = service.files().create(
        body=file_metadata,
        media_body=media,
        fields="id, name, webViewLink",
    ).execute()

    doc_id = created.get("id")
    grant_file_permissions(
        service,
        file_id=doc_id,
        role="writer",
        share_public=True,
        additional_email=additional_email,
    )

    return {
        "doc_id": doc_id,
        "name": created.get("name"),
        "url": created.get("webViewLink"),
    }


def upload_pdf_to_google_drive(
    service,
    pdf_path: Path,
    pdf_name: str,
    parent_folder_id: str,
    additional_email: Optional[str] = None,
) -> Dict[str, str]:
    """Uploads a searchable 2-page English Section PDF into the target Google Drive folder."""
    final_name = get_versioned_filename(service, parent_folder_id, pdf_name, replace_existing=True)
    file_metadata = {
        "name": final_name,
        "parents": [parent_folder_id],
        "mimeType": "application/pdf",
    }
    media = MediaFileUpload(
        str(pdf_path),
        mimetype="application/pdf",
        resumable=True,
    )
    created = service.files().create(
        body=file_metadata,
        media_body=media,
        fields="id, name, webViewLink",
    ).execute()
    pdf_id = created.get("id")
    grant_file_permissions(
        service,
        file_id=pdf_id,
        role="reader",
        share_public=True,
        additional_email=additional_email,
    )
    return {
        "file_id": pdf_id,
        "name": created.get("name"),
        "url": created.get("webViewLink"),
    }


# ─────────────────────────────────────────────────────────────────────────────
# 6. HIGH-LEVEL SECTION & CHAPTER PACKAGE EXPORTERS
# ─────────────────────────────────────────────────────────────────────────────

def export_section_packet_to_google_drive(
    section_meta: Dict[str, Any],
    packet_data: Dict[str, Any],
    include_scaffolded: bool = True,
    include_section_pdf: bool = True,
    additional_email: Optional[str] = None,
    service=None,
) -> Dict[str, Any]:
    """
    Exports a single section's Student Worksheet (Standard + optional Scaffolded),
    Section Teacher Plan & Clipboard Key, and 2-page English Textbook Section PDF into:
      #US History (English) Curriculum Materials / Chp XX - [Title] / Sec. X.X - [Title]
    """
    if service is None:
        service = get_gdrive_service()

    ch_num = section_meta["chapter_num"]
    ch_title = section_meta["chapter_title"]
    sec_code = section_meta["section_code"]
    sec_title = section_meta["title"]

    section_folder_id, folder_path_display, section_folder_url = get_target_section_folder(
        service,
        chapter_num=ch_num,
        chapter_title=ch_title,
        section_code=sec_code,
        section_title=sec_title,
        additional_email=additional_email,
    )

    TEMP_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    base_prefix = f"Chp {ch_num:02d} - Sec. {sec_code}"

    # 1. Standard Student Packet
    std_base = f"{base_prefix} - Student Worksheet (Standard)"
    std_doc_name = get_versioned_filename(service, section_folder_id, std_base)
    local_std_path = TEMP_DOCS_DIR / f"{std_doc_name}.docx"
    build_docx_file(
        markdown_text=packet_data["student_standard_md"],
        doc_title=std_doc_name,
        is_student_handout=True,
        output_path=local_std_path,
        compact_mode=True,
        omit_auto_title=True,
        include_top_student_header=True,
    )
    std_res = upload_docx_as_google_doc(
        service,
        docx_path=local_std_path,
        doc_name=std_doc_name,
        parent_folder_id=section_folder_id,
        additional_email=additional_email,
    )

    # 2. Optional Scaffolded Student Packet
    scaf_res = None
    if include_scaffolded and packet_data.get("student_scaffolded_md"):
        scaf_base = f"{base_prefix} - Student Worksheet (Scaffolded Support)"
        scaf_doc_name = get_versioned_filename(service, section_folder_id, scaf_base)
        local_scaf_path = TEMP_DOCS_DIR / f"{scaf_doc_name}.docx"
        build_docx_file(
            markdown_text=packet_data["student_scaffolded_md"],
            doc_title=scaf_doc_name,
            is_student_handout=True,
            output_path=local_scaf_path,
            compact_mode=True,
            omit_auto_title=True,
            include_top_student_header=True,
        )
        scaf_res = upload_docx_as_google_doc(
            service,
            docx_path=local_scaf_path,
            doc_name=scaf_doc_name,
            parent_folder_id=section_folder_id,
            additional_email=additional_email,
        )

    # 3. Section Teacher Plan & Clipboard Key
    teacher_base = f"{base_prefix} - Teacher Plan & Clipboard Key"
    teacher_doc_name = get_versioned_filename(service, section_folder_id, teacher_base)
    local_teacher_path = TEMP_DOCS_DIR / f"{teacher_doc_name}.docx"
    build_docx_file(
        markdown_text=packet_data["teacher_key_md"],
        doc_title=teacher_doc_name,
        is_student_handout=False,
        output_path=local_teacher_path,
        compact_mode=True,
        omit_auto_title=True,
        include_top_student_header=False,
    )
    teacher_res = upload_docx_as_google_doc(
        service,
        docx_path=local_teacher_path,
        doc_name=teacher_doc_name,
        parent_folder_id=section_folder_id,
        additional_email=additional_email,
    )

    # 4. Searchable 2-Page English Section PDF
    pdf_res = None
    if include_section_pdf and section_meta.get("pdf_file"):
        pdf_raw = str(section_meta["pdf_file"]).replace("\\", "/")
        pdf_p = Path(pdf_raw)
        if not pdf_p.exists():
            root_pdf = Path(__file__).resolve().parent.parent / pdf_raw
            if root_pdf.exists():
                pdf_p = root_pdf
        if pdf_p.exists():
            pdf_name = f"Sec {sec_code} - {sec_title} (Textbook pp. {section_meta.get('printed_pages', '')}).pdf"
            pdf_res = upload_pdf_to_google_drive(
                service,
                pdf_path=pdf_p,
                pdf_name=pdf_name,
                parent_folder_id=section_folder_id,
                additional_email=additional_email,
            )

    return {
        "folder_path": folder_path_display,
        "folder_id": section_folder_id,
        "folder_url": section_folder_url,
        "student_standard": std_res,
        "student_scaffolded": scaf_res,
        "teacher_key": teacher_res,
        "textbook_pdf": pdf_res,
    }


def export_chapter_package_to_google_drive(
    chapter_num: int,
    chapter_title: str,
    master_plan_md: str,
    standard_booklet_md: str,
    scaffolded_booklet_md: Optional[str] = None,
    assessment_data: Optional[Dict[str, str]] = None,
    section_items_with_packets: Optional[List[Tuple[Dict[str, Any], Dict[str, Any]]]] = None,
    additional_email: Optional[str] = None,
    progress_callback=None,
) -> Dict[str, Any]:
    """
    Exports the complete English Chapter Package to:
      #US History (English) Curriculum Materials / Chp XX - [Title]
    Includes:
      1. Chapter Master Teacher Plan & All-in-One Clipboard Answer Key
      2. Chapter Student Packet Booklet (Standard Edition)
      3. Chapter Student Packet Booklet (Scaffolded Support Edition, if enabled)
      4. End-of-Chapter Central Idea Review & Assessment + Answer Key (if provided)
      5. Individual Section Subfolders with per-section Google Docs & 2-Page Textbook PDFs (if provided)
    """
    service = get_gdrive_service()
    chapter_folder_id, chapter_path_display, chapter_folder_url = get_target_chapter_folder(
        service,
        chapter_num=chapter_num,
        chapter_title=chapter_title,
        additional_email=additional_email,
    )

    TEMP_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    ch_prefix = f"Chp {chapter_num:02d}"

    if progress_callback:
        progress_callback("Uploading Master Teacher Plan & All-in-One Clipboard Key...")

    # 1. Master Teacher Plan & Clipboard Key
    master_base = f"{ch_prefix} - Master Teacher Plan & Clipboard Key"
    master_doc_name = get_versioned_filename(service, chapter_folder_id, master_base)
    local_master_path = TEMP_DOCS_DIR / f"{master_doc_name}.docx"
    build_docx_file(
        markdown_text=master_plan_md,
        doc_title=master_doc_name,
        is_student_handout=False,
        output_path=local_master_path,
        compact_mode=True,
        omit_auto_title=True,
        include_top_student_header=False,
    )
    master_res = upload_docx_as_google_doc(
        service,
        docx_path=local_master_path,
        doc_name=master_doc_name,
        parent_folder_id=chapter_folder_id,
        additional_email=additional_email,
    )

    if progress_callback:
        progress_callback("Uploading Chapter Student Packet Booklet (Standard Edition)...")

    # 2. Standard Student Packet Booklet
    std_base = f"{ch_prefix} - Student Packet Booklet (Standard Edition)"
    std_doc_name = get_versioned_filename(service, chapter_folder_id, std_base)
    local_std_path = TEMP_DOCS_DIR / f"{std_doc_name}.docx"
    build_docx_file(
        markdown_text=standard_booklet_md,
        doc_title=std_doc_name,
        is_student_handout=True,
        output_path=local_std_path,
        compact_mode=True,
        omit_auto_title=True,
        include_top_student_header=True,
    )
    std_booklet_res = upload_docx_as_google_doc(
        service,
        docx_path=local_std_path,
        doc_name=std_doc_name,
        parent_folder_id=chapter_folder_id,
        additional_email=additional_email,
    )

    # 3. Scaffolded Student Packet Booklet (Optional)
    scaf_booklet_res = None
    if scaffolded_booklet_md:
        if progress_callback:
            progress_callback("Uploading Chapter Student Packet Booklet (Scaffolded Support Edition)...")
        scaf_base = f"{ch_prefix} - Student Packet Booklet (Scaffolded Support Edition)"
        scaf_doc_name = get_versioned_filename(service, chapter_folder_id, scaf_base)
        local_scaf_path = TEMP_DOCS_DIR / f"{scaf_doc_name}.docx"
        build_docx_file(
            markdown_text=scaffolded_booklet_md,
            doc_title=scaf_doc_name,
            is_student_handout=True,
            output_path=local_scaf_path,
            compact_mode=True,
            omit_auto_title=True,
            include_top_student_header=True,
        )
        scaf_booklet_res = upload_docx_as_google_doc(
            service,
            docx_path=local_scaf_path,
            doc_name=scaf_doc_name,
            parent_folder_id=chapter_folder_id,
            additional_email=additional_email,
        )

    # 4. End-of-Chapter Central Idea Review & Assessment (Optional)
    assessment_res = {}
    if assessment_data:
        if progress_callback:
            progress_callback("Uploading End-of-Chapter Central Idea Review & Assessment...")
        for key_label, doc_suffix, is_stud in [
            ("standard_test_md", "Central Idea Review & Assessment (Standard)", True),
            ("scaffolded_test_md", "Central Idea Review & Assessment (Scaffolded)", True),
            ("teacher_key_md", "Central Idea Review & Assessment (Teacher Answer Key)", False),
        ]:
            md_txt = assessment_data.get(key_label)
            if md_txt:
                t_base = f"{ch_prefix} - {doc_suffix}"
                t_name = get_versioned_filename(service, chapter_folder_id, t_base)
                t_path = TEMP_DOCS_DIR / f"{t_name}.docx"
                build_docx_file(
                    markdown_text=md_txt,
                    doc_title=t_name,
                    is_student_handout=is_stud,
                    output_path=t_path,
                    compact_mode=True,
                    omit_auto_title=True,
                    include_top_student_header=is_stud,
                )
                assessment_res[key_label] = upload_docx_as_google_doc(
                    service,
                    docx_path=t_path,
                    doc_name=t_name,
                    parent_folder_id=chapter_folder_id,
                    additional_email=additional_email,
                )

    # 5. Individual Section Subfolders (Optional)
    section_exports = {}
    if section_items_with_packets:
        total_secs = len(section_items_with_packets)
        for idx, (sec_obj, pkt_data) in enumerate(section_items_with_packets, start=1):
            scode = str(sec_obj["section_code"])
            if progress_callback:
                progress_callback(f"Uploading Section {scode} Handouts & 2-Page PDF ({idx}/{total_secs})...")
            sec_res = export_section_packet_to_google_drive(
                section_meta=sec_obj,
                packet_data=pkt_data,
                include_scaffolded=bool(scaffolded_booklet_md),
                include_section_pdf=True,
                additional_email=additional_email,
                service=service,
            )
            section_exports[scode] = sec_res

    return {
        "chapter_folder_path": chapter_path_display,
        "chapter_folder_id": chapter_folder_id,
        "chapter_folder_url": chapter_folder_url,
        "master_teacher_plan": master_res,
        "standard_student_booklet": std_booklet_res,
        "scaffolded_student_booklet": scaf_booklet_res,
        "assessment": assessment_res,
        "sections": section_exports,
    }


def export_assessment_package_to_google_drive(
    chapter_num: int,
    chapter_title: str,
    assessment_data: Dict[str, str],
    additional_email: Optional[str] = None,
    service=None,
) -> Dict[str, Any]:
    """
    Exports the End-of-Chapter Central Idea Review & Assessment Google Docs
    (Standard, Scaffolded, Teacher Key) into the Chapter's Google Drive folder.
    """
    if service is None:
        service = get_gdrive_service()

    chapter_folder_id, chapter_path_display, chapter_folder_url = get_target_chapter_folder(
        service,
        chapter_num=chapter_num,
        chapter_title=chapter_title,
        additional_email=additional_email,
    )

    TEMP_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    ch_prefix = f"Chp {chapter_num:02d}"

    docs_res = {}
    for key_label, doc_suffix, is_stud in [
        ("standard_test_md", "Central Idea Review & Assessment (Standard)", True),
        ("scaffolded_test_md", "Central Idea Review & Assessment (Scaffolded)", True),
        ("teacher_key_md", "Central Idea Review & Assessment (Teacher Answer Key)", False),
    ]:
        md_txt = assessment_data.get(key_label)
        if md_txt:
            t_base = f"{ch_prefix} - {doc_suffix}"
            t_name = get_versioned_filename(service, chapter_folder_id, t_base)
            t_path = TEMP_DOCS_DIR / f"{t_name}.docx"
            build_docx_file(
                markdown_text=md_txt,
                doc_title=t_name,
                is_student_handout=is_stud,
                output_path=t_path,
                compact_mode=True,
                omit_auto_title=True,
                include_top_student_header=is_stud,
            )
            docs_res[key_label] = upload_docx_as_google_doc(
                service,
                docx_path=t_path,
                doc_name=t_name,
                parent_folder_id=chapter_folder_id,
                additional_email=additional_email,
            )

    return {
        "chapter_folder_path": chapter_path_display,
        "chapter_folder_id": chapter_folder_id,
        "chapter_folder_url": chapter_folder_url,
        "docs": docs_res,
    }


def get_docx_bytes(
    markdown_text: str,
    filename_stem: str,
    is_student_handout: bool,
    include_top_student_header: bool = True,
) -> bytes:
    """
    Generates and returns in-memory DOCX bytes from markdown content.
    Uses standard layout rules (compact spacing, native tables, metadata pills).
    """
    TEMP_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TEMP_DOCS_DIR / f"{filename_stem}.docx"
    build_docx_file(
        markdown_text=markdown_text,
        doc_title=filename_stem,
        is_student_handout=is_student_handout,
        output_path=out_path,
        compact_mode=True,
        omit_auto_title=True,
        include_top_student_header=include_top_student_header,
    )
    return out_path.read_bytes()

